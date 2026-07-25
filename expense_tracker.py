"""Local BofA statement importer and compact Excel expense reporter."""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import sqlite3
import subprocess
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
OUTPUT_DIR = APP_DIR / "output"
DATABASE = DATA_DIR / "tracker.sqlite3"
WORKBOOK = OUTPUT_DIR / "expense_tracker.xlsx"
PARSER_VERSION = "3"  # Includes full BofA row/reference and the credit/debit source type.

CATEGORIES = [
    "Groceries", "Beauty", "Car parts", "Home cleaning", "House bills",
    "Restaurants", "Shopping", "Online shopping", "Entertainment & subscriptions", "Insurance",
    "Fuel", "Pet care", "Alcohol", "Healthcare", "Home improvement", "Dealership", "Having fun", "Other",
]
DEFAULT_RULES = {
    "WAL MART": ("Groceries", "Walmart"), "WALMART": ("Groceries", "Walmart"),
    "COSTCO": ("Groceries", "Costco"), "HOLIDAY MARKET": ("Groceries", "Holiday Market"),
    "ASIAN MART": ("Groceries", "Asian Mart"), "KROGER": ("Groceries", "Kroger"),
    "MEIJER": ("Groceries", "Meijer"), "ALDI": ("Groceries", "Aldi"),
    "TRADER JOE": ("Groceries", "Trader Joe's"),
    "PANERA": ("Restaurants", "Panera Bread"), "MCDONALD": ("Restaurants", "McDonald's"),
    "STARBUCKS": ("Restaurants", "Starbucks"), "DOORDASH": ("Restaurants", "DoorDash"),
    "UBER EATS": ("Restaurants", "Uber Eats"), "RESTAURANT": ("Restaurants", "Restaurant"),
    "COSTCO GAS": ("Fuel", "Costco Gas"), "SHELL": ("Fuel", "Shell"), "BP ": ("Fuel", "BP"),
    "EXXON": ("Fuel", "Exxon"), "MOBIL": ("Fuel", "Mobil"),
    "CONSUMERS ENERGY": ("House bills", "Gas / Electricity"), "DTE ENERGY": ("House bills", "Gas / Electricity"),
    "ATT": ("House bills", "Phone / Internet"), "VERIZON": ("House bills", "Phone / Internet"),
    "METLIFE": ("Insurance", "MetLife"), "STATE FARM": ("Insurance", "State Farm"),
    "ALLSTATE": ("Insurance", "Allstate"), "GEICO": ("Insurance", "Geico"),
    "HAMILTON ANIMAL": ("Pet care", "Veterinary"), "VET": ("Pet care", "Veterinary"),
    "MICHAELS WINE": ("Alcohol", "Michaels Wine Shoppe"),
}
INTERNAL_TERMS = (
    "PAYMENT FROM CHK", "MOBILE BANKING PAYMENT TO CRD", "TRANSFER TO SAV",
    "TRANSFER TO CHK", "ONLINE BANKING TRANSFER TO",
)
AMBIGUOUS_TERMS = (
    "ZELLE", "VENMO", "CASH APP", "WITHDRWL", "ATM", "COINBASE", "CRYPTO",
    "AMAZON", "GOOGLE", "PAYMENT", "TRANSFER",
)

# Most statements expose only a merchant string, not a standard industry code.
# These high-confidence patterns make useful local suggestions without sending
# private transaction information to an external service.
SMART_RULES: tuple[tuple[str, str, str], ...] = (
    ("COSTCO GAS", "Fuel", "Costco Gas"),
    ("NEW YORK INTERNATIONAL", "Groceries", "New York International"),
    ("WM SUPERCENTER", "Groceries", "Walmart Supercenter"),
    ("SHELL", "Fuel", "Shell"), ("EXXON", "Fuel", "Exxon"), ("MOBIL", "Fuel", "Mobil"),
    ("PANERA", "Restaurants", "Panera Bread"), ("MCDONALD", "Restaurants", "McDonald's"),
    ("STARBUCKS", "Restaurants", "Starbucks"), ("DOORDASH", "Restaurants", "DoorDash"),
    ("UBER EATS", "Restaurants", "Uber Eats"), ("GRUBHUB", "Restaurants", "Grubhub"),
    ("TOUS LES JOURS", "Restaurants", "Tous Les Jours"),
    ("EARLY BIRD CAFE", "Restaurants", "Early Bird Cafe"),
    ("OLIVE GARDEN", "Restaurants", "Olive Garden"),
    ("OSAKA", "Restaurants", "Osaka Japanese Steakhouse"), ("RESTAURANT", "Restaurants", "Restaurant"),
    ("AMAZON", "Online shopping", "Amazon"), ("AMZN", "Online shopping", "Amazon"),
    ("ALIEXPRESS", "Online shopping", "AliExpress"), ("EBAY", "Online shopping", "eBay"),
    ("ETSY", "Online shopping", "Etsy"), ("WALMART COM", "Online shopping", "Walmart.com"),
    ("SOMA INTIMATES", "Shopping", "Soma Intimates"),
    ("YOUTUBE", "Entertainment & subscriptions", "YouTube"),
    ("NETFLIX", "Entertainment & subscriptions", "Netflix"),
    ("SPOTIFY", "Entertainment & subscriptions", "Spotify"), ("HULU", "Entertainment & subscriptions", "Hulu"),
    ("APPLE COM BILL", "Entertainment & subscriptions", "Apple services"),
    ("HAMILTON ANIMAL", "Pet care", "Veterinary"), ("VET", "Pet care", "Veterinary"),
    ("ART DENTISTRY", "Healthcare", "Dentistry"), ("DENTIST", "Healthcare", "Dentistry"),
    ("HOME DEPOT", "Home improvement", "Home Depot"), ("LOWE S", "Home improvement", "Lowe's"),
    ("HALL OF FAME BILLIARDS", "Having fun", "Hall Of Fame Billiards"),
    ("SUBURBAN FORD", "Dealership", "Suburban Ford Of Troy"),
    ("JOERANDAZZOSFRUIT", "Groceries", "Joe Randazzo's Fruit & Vegetables"),
)

# These are intentionally broader than SMART_RULES and run only after exact
# merchant and user-saved rules. They give a useful category to new local
# merchants while retaining the merchant name as the report subcategory.
KEYWORD_CATEGORIES: tuple[tuple[tuple[str, ...], str], ...] = (
    (("GROCERY", "SUPERMARKET", "FRUIT", "VEG", "PRODUCE", "MARKET", "FOODS"), "Groceries"),
    (("CAFE", "GRILL", "PIZZA", "BURGER", "KITCHEN", "DINER", "BAKERY", "BISTRO", "SUSHI", "STEAKHOUSE"), "Restaurants"),
    (("BOUTIQUE", "APPAREL", "CLOTHING", "OUTLET", "SHOES", "JEWELRY"), "Shopping"),
)


class ImportErrorWithHelp(ValueError):
    pass


@dataclass
class Transaction:
    account: str
    month: str
    posted_date: str
    description: str
    amount: Decimal
    fingerprint: str
    source_type: str = "unknown"
    category: str | None = None
    subcategory: str | None = None


def clean_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def money(value: str) -> Decimal:
    try:
        return Decimal(value.replace("$", "").replace(",", ""))
    except InvalidOperation as exc:
        raise ImportErrorWithHelp(f"Could not read amount {value!r}.") from exc


def normalized_merchant(description: str) -> str:
    text = description.upper()
    text = re.sub(r"\b\d{2}/\d{2}(?:/\d{2})?\b", " ", text)
    text = re.sub(r"#?\d{4,}\b", " ", text)
    text = re.sub(r"\b(PURCHASE|CHECKCARD|WITHDRWL|DES|ID|WEB|CONF)\b", " ", text)
    text = re.sub(r"[^A-Z&* ]", " ", text)
    return clean_spaces(text)[:100]


def suggested_rule(merchant: str) -> tuple[str, str] | None:
    """Return a high-confidence local category suggestion for a merchant."""
    for pattern, category, subcategory in SMART_RULES:
        if pattern in merchant:
            return category, subcategory
    for key, rule in DEFAULT_RULES.items():
        if key in merchant:
            return rule
    for keywords, category in KEYWORD_CATEGORIES:
        if any(keyword in merchant for keyword in keywords):
            return category, merchant.title()
    return None


def extract_text(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ImportErrorWithHelp(f"{path.name} is password-protected. Save an unlocked copy first.")
        text = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in reader.pages)
    except ImportErrorWithHelp:
        raise
    except Exception as exc:
        raise ImportErrorWithHelp(f"Could not open {path.name}: {exc}") from exc
    if len(re.sub(r"\s+", "", text)) < 200:
        raise ImportErrorWithHelp(f"{path.name} has no readable text. Scanned/image-only PDFs are not supported.")
    return text


def statement_info(text: str) -> tuple[str, str, str]:
    account_match = re.search(r"Account(?:\s+Number:|\s*#)\s*([0-9 ]{8,})", text, re.I)
    if not account_match:
        raise ImportErrorWithHelp("This does not look like a supported Bank of America statement (account number not found).")
    account = re.sub(r"\D", "", account_match.group(1))[-4:]
    credit = re.search(r"!\s*([A-Z][a-z]+)\s+(\d{1,2})\s*-\s*([A-Z][a-z]+)\s+(\d{1,2}),\s*(\d{4})", text)
    debit = re.search(r"for\s+([A-Z][a-z]+)\s+(\d{1,2}),\s*(\d{4})\s+to\s+([A-Z][a-z]+)\s+(\d{1,2}),\s*(\d{4})", text, re.I)
    if credit:
        closing = datetime.strptime(f"{credit.group(3)} {credit.group(4)} {credit.group(5)}", "%B %d %Y")
        return account, closing.strftime("%Y-%m"), "credit"
    if debit:
        closing = datetime.strptime(f"{debit.group(4)} {debit.group(5)} {debit.group(6)}", "%B %d %Y")
        return account, closing.strftime("%Y-%m"), "debit"
    raise ImportErrorWithHelp("Could not find the statement period in this BofA PDF.")


def section(text: str, start: str, end: str) -> str:
    found = re.search(start + r"(.*?)" + end, text, re.I | re.S)
    return found.group(1) if found else ""


def parse_credit(text: str, account: str, month: str) -> list[Transaction]:
    body = section(text, r"Transactions\s+.*?Purchases and Adjustments", r"TOTAL PURCHASES AND ADJUSTMENTS")
    if not body:
        raise ImportErrorWithHelp("Could not find the credit-card purchase table.")
    return parse_rows(body, account, month, allow_positive=True, statement_type="credit")


def parse_debit(text: str, account: str, month: str) -> list[Transaction]:
    # Debit statements have the card table followed by other withdrawals; both are expenses except internal transfers/cards.
    body = section(text, r"ATM and debit card subtractions(?:\s*- continued)?", r"Total other subtractions|Service fees")
    if not body:
        raise ImportErrorWithHelp("Could not find the debit-account withdrawal tables.")
    return parse_rows(body, account, month, allow_positive=False, statement_type="debit")


def parse_rows(body: str, account: str, month: str, allow_positive: bool, statement_type: str) -> list[Transaction]:
    rows: list[Transaction] = []
    lines = [clean_spaces(line) for line in body.splitlines() if clean_spaces(line)]
    # BofA's extracted rows begin with date; the final number on the same line is its amount.
    pattern = re.compile(r"^(\d{2}/\d{2}(?:/\d{2})?)\s+(.*?)(-?\$?[\d,]+\.\d{2})$")
    pending = ""
    for line in lines:
        upper_line = line.upper()
        if (upper_line.startswith(("DATE ", "DESCRIPTION", "AMOUNT", "TOTAL", "CONTINUED"))
                or "WITHDRAWALS AND OTHER SUBTRACTIONS" in upper_line
                or "ATM AND DEBIT CARD SUBTRACTIONS" in upper_line
                or upper_line == "OTHER SUBTRACTIONS"):
            continue
        match = pattern.match(line)
        if not match:
            if re.match(r"^[A-Z0-9*#& .,'/-]+$", line, re.I) and not re.match(r"^\d{4}$", line):
                pending = (pending + " " + line).strip()
            continue
        date_text, description, amount_text = match.groups()
        description = clean_spaces((pending + " " + description).strip())
        pending = ""
        amount = money(amount_text)
        if statement_type == "credit" and amount <= 0:
            continue
        if statement_type == "debit" and amount >= 0:
            continue
        if any(term in description.upper() for term in INTERNAL_TERMS):
            continue
        # Debit amounts are negative on the statement; reports use positive expense values.
        expense = abs(amount)
        merchant = normalized_merchant(description)
        # Same merchant/date/amount can legitimately occur more than once. BofA
        # includes a distinct reference number in the full description, so retain
        # that complete row when building the de-duplication identity.
        source_row = clean_spaces(description).upper()
        fingerprint = hashlib.sha256(f"{account}|{month}|{date_text}|{source_row}|{expense}".encode()).hexdigest()
        rows.append(Transaction(account, month, date_text, description, expense, fingerprint, statement_type))
    if not rows:
        raise ImportErrorWithHelp("No transaction rows could be read. The BofA PDF layout may have changed.")
    return rows


def parse_statement(path: Path) -> tuple[str, str, str, list[Transaction], str]:
    text = extract_text(path)
    account, month, statement_type = statement_info(text)
    transactions = parse_credit(text, account, month) if statement_type == "credit" else parse_debit(text, account, month)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return account, month, statement_type, transactions, digest


class Store:
    def __init__(self) -> None:
        DATA_DIR.mkdir(exist_ok=True)
        self.connection = sqlite3.connect(DATABASE)
        self.connection.row_factory = sqlite3.Row
        self.connection.executescript("""
            CREATE TABLE IF NOT EXISTS statements (
                digest TEXT PRIMARY KEY, account TEXT NOT NULL, month TEXT NOT NULL, source_name TEXT NOT NULL,
                parser_version TEXT NOT NULL DEFAULT '1'
            );
            CREATE TABLE IF NOT EXISTS transactions (
                fingerprint TEXT PRIMARY KEY, statement_digest TEXT NOT NULL, account TEXT NOT NULL, month TEXT NOT NULL,
                posted_date TEXT NOT NULL, description TEXT NOT NULL, merchant TEXT NOT NULL, amount TEXT NOT NULL,
                source_type TEXT NOT NULL DEFAULT 'unknown',
                FOREIGN KEY(statement_digest) REFERENCES statements(digest)
            );
            CREATE TABLE IF NOT EXISTS merchant_rules (
                merchant TEXT PRIMARY KEY, category TEXT NOT NULL, subcategory TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS categories (
                name TEXT PRIMARY KEY
            );
        """)
        statement_columns = {row["name"] for row in self.connection.execute("PRAGMA table_info(statements)")}
        if "parser_version" not in statement_columns:
            self.connection.execute("ALTER TABLE statements ADD COLUMN parser_version TEXT NOT NULL DEFAULT '1'")
        transaction_columns = {row["name"] for row in self.connection.execute("PRAGMA table_info(transactions)")}
        if "source_type" not in transaction_columns:
            self.connection.execute("ALTER TABLE transactions ADD COLUMN source_type TEXT NOT NULL DEFAULT 'unknown'")
        self.connection.commit()

    def rule_for(self, merchant: str) -> tuple[str, str] | None:
        row = self.connection.execute("SELECT category, subcategory FROM merchant_rules WHERE merchant = ?", (merchant,)).fetchone()
        if row:
            return row["category"], row["subcategory"]
        return suggested_rule(merchant)

    def save_rule(self, merchant: str, category: str, subcategory: str) -> None:
        self.connection.execute("INSERT OR IGNORE INTO categories(name) VALUES (?)", (category,))
        self.connection.execute("INSERT INTO merchant_rules(merchant, category, subcategory) VALUES (?, ?, ?) ON CONFLICT(merchant) DO UPDATE SET category=excluded.category, subcategory=excluded.subcategory", (merchant, category, subcategory))
        self.connection.commit()

    def categories(self) -> list[str]:
        custom = [row["name"] for row in self.connection.execute("SELECT name FROM categories ORDER BY name")]
        return sorted(set(CATEGORIES).union(custom))

    def other_merchants(self) -> list[tuple[str, str, int, Decimal]]:
        """Merchants that currently fall back to Other, grouped for the edit dialog."""
        rows = self.connection.execute(
            "SELECT merchant, MIN(description) AS description, COUNT(*) AS count, amount FROM transactions GROUP BY merchant"
        ).fetchall()
        result: list[tuple[str, str, int, Decimal]] = []
        for row in rows:
            if self.rule_for(row["merchant"]) is None:
                amount = sum(
                    (Decimal(item["amount"]) for item in self.connection.execute(
                        "SELECT amount FROM transactions WHERE merchant = ?", (row["merchant"],)
                    )),
                    Decimal(),
                )
                result.append((row["merchant"], row["description"], row["count"], amount))
        return sorted(result, key=lambda item: item[0])

    def months(self) -> list[str]:
        return [row["month"] for row in self.connection.execute("SELECT DISTINCT month FROM transactions ORDER BY month")]

    def clear_transaction_history(self) -> None:
        """Remove imported statement data while preserving merchant/category rules."""
        self.connection.execute("DELETE FROM transactions")
        self.connection.execute("DELETE FROM statements")
        self.connection.commit()

    def save_statement(self, digest: str, source_name: str, account: str, month: str, transactions: list[Transaction]) -> int:
        existing = self.connection.execute("SELECT parser_version FROM statements WHERE digest = ?", (digest,)).fetchone()
        if existing and existing["parser_version"] == PARSER_VERSION:
            return 0
        # A BofA statement is a complete view of one account and closing month.
        # Replacing an amended/re-downloaded statement prevents obsolete rows from
        # remaining in that month's report while preserving other accounts' rows.
        previous = self.connection.execute(
            "SELECT digest FROM statements WHERE account = ? AND month = ?", (account, month)
        ).fetchall()
        for row in previous:
            self.connection.execute("DELETE FROM transactions WHERE statement_digest = ?", (row["digest"],))
        self.connection.execute("DELETE FROM statements WHERE account = ? AND month = ?", (account, month))
        self.connection.execute(
            "INSERT INTO statements(digest, account, month, source_name, parser_version) VALUES (?, ?, ?, ?, ?)",
            (digest, account, month, source_name, PARSER_VERSION),
        )
        count = 0
        for txn in transactions:
            result = self.connection.execute(
                "INSERT OR IGNORE INTO transactions VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (txn.fingerprint, digest, txn.account, txn.month, txn.posted_date, txn.description,
                 normalized_merchant(txn.description), str(txn.amount), txn.source_type),
            )
            count += result.rowcount
        self.connection.commit()
        return count

    def has_transaction(self, fingerprint: str) -> bool:
        return self.connection.execute("SELECT 1 FROM transactions WHERE fingerprint = ?", (fingerprint,)).fetchone() is not None

    def totals_for_month(self, month: str) -> dict[tuple[str, str], Decimal]:
        totals: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
        rows = self.connection.execute("SELECT merchant, amount FROM transactions WHERE month = ?", (month,)).fetchall()
        for row in rows:
            rule = self.rule_for(row["merchant"])
            # Automatic export keeps ordinary new merchants visible without
            # permanently guessing a rule. A later review can refine them.
            category, subcategory = rule or ("Other", row["merchant"].title())
            totals[(category, subcategory)] += Decimal(row["amount"])
        return totals

    def totals_for_month_by_source(self, month: str) -> dict[tuple[str, str, str], Decimal]:
        totals: dict[tuple[str, str, str], Decimal] = defaultdict(Decimal)
        rows = self.connection.execute("SELECT merchant, amount, source_type FROM transactions WHERE month = ?", (month,)).fetchall()
        for row in rows:
            category, subcategory = self.rule_for(row["merchant"]) or ("Other", row["merchant"].title())
            totals[(category, subcategory, row["source_type"])] += Decimal(row["amount"])
        return totals

    def close(self) -> None:
        self.connection.close()


def write_month(store: Store, month: str) -> Path:
    OUTPUT_DIR.mkdir(exist_ok=True)
    workbook = load_workbook(WORKBOOK) if WORKBOOK.exists() else Workbook()
    if month in workbook.sheetnames:
        del workbook[month]
    if workbook.sheetnames == ["Sheet"] and workbook["Sheet"].max_row == 1 and workbook["Sheet"]["A1"].value is None:
        ws = workbook["Sheet"]
        ws.title = month
    else:
        ws = workbook.create_sheet(month)
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"Monthly expenses — {month}"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A1:D1")
    headers = ["Category", "Subcategory", "Payment source", "Total Expense"]
    for column, header in enumerate(headers, 1):
        cell = ws.cell(3, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center")
    totals = store.totals_for_month_by_source(month)
    row = 4
    category_totals: dict[str, Decimal] = defaultdict(Decimal)
    for (category, subcategory, source_type), value in totals.items():
        category_totals[category] += value
    for category in sorted(category_totals):
        ws.cell(row, 1, category).font = Font(bold=True)
        ws.cell(row, 4, category_totals[category]).font = Font(bold=True)
        ws.cell(row, 4).number_format = '$#,##0.00'
        ws.cell(row, 1).fill = ws.cell(row, 2).fill = ws.cell(row, 3).fill = ws.cell(row, 4).fill = PatternFill("solid", fgColor="D9EAF7")
        row += 1
        for (item_category, subcategory, source_type), value in sorted(totals.items()):
            if item_category != category:
                continue
            ws.cell(row, 2, subcategory)
            source_label = {"debit": "Debit card", "credit": "Credit card"}.get(source_type, "Unknown")
            ws.cell(row, 3, source_label)
            ws.cell(row, 4, value).number_format = '$#,##0.00'
            if source_type == "debit":
                for column in range(1, 5):
                    ws.cell(row, column).fill = PatternFill("solid", fgColor="FFF2CC")
            row += 1
    overall = sum(totals.values(), Decimal())
    ws.cell(row + 1, 1, "Overall total").font = Font(bold=True)
    ws.cell(row + 1, 4, overall).font = Font(bold=True)
    ws.cell(row + 1, 4).number_format = '$#,##0.00'
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.freeze_panes = "A4"
    workbook.save(WORKBOOK)
    return WORKBOOK


class ReviewDialog(tk.Toplevel):
    def __init__(self, parent: tk.Tk, store: Store, unresolved: dict[str, list[Transaction]]) -> None:
        super().__init__(parent)
        self.title("Review new merchants")
        self.transient(parent)
        self.grab_set()
        self.store, self.unresolved, self.index = store, list(unresolved.items()), 0
        self.result = False
        self.merchant = tk.StringVar()
        self.category = tk.StringVar(value="Other")
        self.subcategory = tk.StringVar()
        frame = ttk.Frame(self, padding=16)
        frame.grid(sticky="nsew")
        ttk.Label(frame, text="This transaction needs your input", font=("Segoe UI", 11, "bold")).grid(row=0, column=0, columnspan=2, sticky="w")
        ttk.Label(frame, textvariable=self.merchant, wraplength=460).grid(row=1, column=0, columnspan=2, pady=(12, 12), sticky="w")
        ttk.Label(frame, text="Category").grid(row=2, column=0, sticky="w")
        ttk.Combobox(frame, textvariable=self.category, values=CATEGORIES, state="readonly", width=30).grid(row=2, column=1, padx=(12, 0))
        ttk.Label(frame, text="Subcategory / merchant").grid(row=3, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(frame, textvariable=self.subcategory, width=33).grid(row=3, column=1, padx=(12, 0), pady=(10, 0))
        self.next_button = ttk.Button(frame, text="Save and continue", command=self.save)
        self.next_button.grid(row=4, column=1, sticky="e", pady=(16, 0))
        self.protocol("WM_DELETE_WINDOW", self.cancel)
        self.show_current()

    def show_current(self) -> None:
        merchant, txns = self.unresolved[self.index]
        examples = "; ".join(t.description for t in txns[:2])
        self.merchant.set(
            f"{merchant}\n{len(txns)} imported transaction(s): {examples}\n\n"
            f"Proposed automatic fallback: Other → {merchant.title()}. "
            "Choose a more accurate category if you know one."
        )
        self.category.set("Other")
        self.subcategory.set(merchant.title())
        self.next_button.config(text="Save and export" if self.index == len(self.unresolved) - 1 else "Save and continue")

    def save(self) -> None:
        subcategory = self.subcategory.get().strip()
        if not subcategory:
            messagebox.showerror("Subcategory required", "Enter a subcategory or merchant name.", parent=self)
            return
        merchant, _ = self.unresolved[self.index]
        self.store.save_rule(merchant, self.category.get(), subcategory)
        self.index += 1
        if self.index == len(self.unresolved):
            self.result = True
            self.destroy()
        else:
            self.show_current()

    def cancel(self) -> None:
        self.result = False
        self.destroy()


class EditOtherDialog(tk.Toplevel):
    """Batch-edit merchants that would otherwise be exported under Other."""
    def __init__(self, parent: tk.Tk, store: Store) -> None:
        super().__init__(parent)
        self.title("Edit Other categories")
        self.geometry("760x440")
        self.minsize(680, 380)
        self.transient(parent)
        self.grab_set()
        self.store = store
        self.items = store.other_merchants()
        self.pending: dict[str, tuple[str, str]] = {}
        self.current: tuple[str, str, int, Decimal] | None = None
        self.category = tk.StringVar()
        self.subcategory = tk.StringVar()

        frame = ttk.Frame(self, padding=16)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="Edit Other categories", font=("Segoe UI", 12, "bold")).grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(frame, text="Select an uncategorized merchant, assign a category, then save all changes.").grid(row=1, column=0, columnspan=3, sticky="w", pady=(3, 12))

        self.listbox = tk.Listbox(frame, width=48, height=14, exportselection=False)
        self.listbox.grid(row=2, column=0, rowspan=6, sticky="nsew")
        self.listbox.bind("<<ListboxSelect>>", self.select_item)
        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=self.listbox.yview)
        scrollbar.grid(row=2, column=1, rowspan=6, sticky="ns")
        self.listbox.config(yscrollcommand=scrollbar.set)

        right = ttk.Frame(frame)
        right.grid(row=2, column=2, rowspan=6, sticky="nsew", padx=(18, 0))
        self.details = ttk.Label(right, text="", wraplength=310, justify="left")
        self.details.pack(anchor="w", pady=(0, 14))
        ttk.Label(right, text="Category").pack(anchor="w")
        self.category_box = ttk.Combobox(right, textvariable=self.category, values=self.store.categories(), width=32)
        self.category_box.pack(fill="x", pady=(2, 10))
        ttk.Button(right, text="Create new category…", command=self.create_category).pack(anchor="w", pady=(0, 14))
        ttk.Label(right, text="Subcategory / merchant").pack(anchor="w")
        ttk.Entry(right, textvariable=self.subcategory, width=35).pack(fill="x", pady=(2, 14))
        self.assign_button = ttk.Button(right, text="Queue assignment", command=self.queue_assignment)
        self.assign_button.pack(anchor="e")

        buttons = ttk.Frame(frame)
        buttons.grid(row=8, column=0, columnspan=3, sticky="e", pady=(16, 0))
        ttk.Button(buttons, text="Cancel", command=self.destroy).pack(side="right")
        ttk.Button(buttons, text="Save changes and close", command=self.save_and_close).pack(side="right", padx=(0, 8))
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(2, weight=1)
        self.refresh_list()

    def refresh_list(self) -> None:
        self.listbox.delete(0, tk.END)
        for merchant, _, count, amount in self.items:
            self.listbox.insert(tk.END, f"{merchant.title()}  ({count} transaction(s), ${amount:,.2f})")
        if self.items:
            self.listbox.selection_set(0)
            self.select_item()
        else:
            self.details.config(text="There are no uncategorized merchants to edit.")
            self.assign_button.config(state="disabled")

    def select_item(self, _event: object | None = None) -> None:
        selected = self.listbox.curselection()
        if not selected:
            return
        self.current = self.items[selected[0]]
        merchant, description, count, amount = self.current
        self.details.config(text=f"{description}\n\n{count} transaction(s), total ${amount:,.2f}")
        self.category.set("Other")
        self.subcategory.set(merchant.title())

    def create_category(self) -> None:
        name = simpledialog.askstring("Create category", "New category name:", parent=self)
        if not name:
            return
        name = clean_spaces(name)
        if not name:
            return
        values = sorted(set(self.category_box.cget("values")).union({name}))
        self.category_box.config(values=values)
        self.category.set(name)

    def queue_assignment(self) -> None:
        if self.current is None:
            return
        category, subcategory = clean_spaces(self.category.get()), clean_spaces(self.subcategory.get())
        if not category or not subcategory:
            messagebox.showerror("Category required", "Enter both a category and a subcategory.", parent=self)
            return
        self.pending[self.current[0]] = (category, subcategory)
        self.items.pop(self.listbox.curselection()[0])
        self.current = None
        self.refresh_list()

    def save_and_close(self) -> None:
        for merchant, (category, subcategory) in self.pending.items():
            self.store.save_rule(merchant, category, subcategory)
        if self.pending:
            messagebox.showinfo("Categories saved", f"Saved {len(self.pending)} categorization(s). Use Export all saved months to update Excel.", parent=self)
        self.destroy()


class App(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("BofA Monthly Expense Tracker")
        self.geometry("800x460")
        self.minsize(720, 400)
        self.store = Store()
        self.paths: list[Path] = []
        self.status = tk.StringVar(value="Select one or more Bank of America statement PDFs.")
        frame = ttk.Frame(self, padding=20)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="Bank of America Monthly Expense Tracker", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        ttk.Label(frame, text="Imports local PDF statements and updates a compact Excel summary.").pack(anchor="w", pady=(4, 14))
        controls = ttk.Frame(frame)
        controls.pack(fill="x")
        first_row = ttk.Frame(controls)
        first_row.pack(fill="x")
        ttk.Button(first_row, text="Select PDF statements", command=self.choose_files).pack(side="left")
        ttk.Button(first_row, text="Export automatically", command=self.export_automatically).pack(side="left", padx=8)
        ttk.Button(first_row, text="Export all saved months", command=self.export_all_months).pack(side="left")
        ttk.Button(first_row, text="Reset ALL…", command=self.reset_all).pack(side="right")
        ttk.Button(first_row, text="Reset transaction history…", command=self.reset_transaction_history).pack(side="right", padx=(0, 8))
        second_row = ttk.Frame(controls)
        second_row.pack(fill="x", pady=(8, 0))
        ttk.Button(second_row, text="Edit Other categories…", command=self.edit_other_categories).pack(side="left")
        ttk.Button(second_row, text="Review all new merchants", command=self.review_and_export).pack(side="left", padx=8)
        self.listbox = tk.Listbox(frame, height=10)
        self.listbox.pack(fill="both", expand=True, pady=12)
        ttk.Label(frame, textvariable=self.status, wraplength=580).pack(anchor="w")
        ttk.Button(frame, text="Open output folder", command=self.open_output).pack(anchor="e", pady=(10, 0))
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def choose_files(self) -> None:
        choices = filedialog.askopenfilenames(title="Select BofA statements", filetypes=[("PDF files", "*.pdf")])
        if choices:
            self.paths = [Path(item) for item in choices]
            self.listbox.delete(0, tk.END)
            for path in self.paths:
                self.listbox.insert(tk.END, path.name)
            self.status.set(f"{len(self.paths)} PDF(s) selected.")

    def export_automatically(self) -> None:
        self.import_files(review_all=False)

    def review_and_export(self) -> None:
        self.import_files(review_all=True)

    def edit_other_categories(self) -> None:
        dialog = EditOtherDialog(self, self.store)
        self.wait_window(dialog)

    def export_all_months(self) -> None:
        months = self.store.months()
        if not months:
            messagebox.showinfo("Nothing to export", "Import at least one statement before exporting saved months.", parent=self)
            return
        for month in months:
            write_month(self.store, month)
        self.status.set(f"Rebuilt {WORKBOOK.name} for: {', '.join(months)}")
        messagebox.showinfo("Export complete", f"Updated {WORKBOOK.name} for: {', '.join(months)}", parent=self)

    @staticmethod
    def requires_review(transaction: Transaction) -> bool:
        description = transaction.description.upper()
        return not normalized_merchant(transaction.description) or any(term in description for term in AMBIGUOUS_TERMS)

    def import_files(self, review_all: bool) -> None:
        if not self.paths:
            messagebox.showinfo("Select statements", "Choose at least one PDF statement first.", parent=self)
            return
        parsed: list[tuple[Path, str, str, list[Transaction], str]] = []
        errors: list[str] = []
        for path in self.paths:
            try:
                account, month, _, transactions, digest = parse_statement(path)
                parsed.append((path, account, month, transactions, digest))
            except ImportErrorWithHelp as exc:
                errors.append(str(exc))
        if errors:
            messagebox.showerror("Nothing was changed", "\n\n".join(errors), parent=self)
            return
        unresolved: dict[str, list[Transaction]] = defaultdict(list)
        for _, _, _, transactions, _ in parsed:
            for txn in transactions:
                merchant = normalized_merchant(txn.description)
                if (not self.store.rule_for(merchant)
                        and not self.store.has_transaction(txn.fingerprint)
                        and (review_all or self.requires_review(txn))):
                    unresolved[merchant].append(txn)
        if unresolved:
            dialog = ReviewDialog(self, self.store, unresolved)
            self.wait_window(dialog)
            if not dialog.result:
                self.status.set("Import cancelled during merchant review. No statements were saved.")
                return
        months: set[str] = set()
        imported = 0
        for path, account, month, transactions, digest in parsed:
            imported += self.store.save_statement(digest, path.name, account, month, transactions)
            months.add(month)
        for month in sorted(months):
            write_month(self.store, month)
        self.status.set(f"Imported {imported} new transaction(s). Workbook updated: {WORKBOOK}")
        messagebox.showinfo("Export complete", f"Updated {WORKBOOK.name} for: {', '.join(sorted(months))}", parent=self)

    def open_output(self) -> None:
        OUTPUT_DIR.mkdir(exist_ok=True)
        if sys.platform.startswith("win"):
            os.startfile(OUTPUT_DIR)  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.run(["open", str(OUTPUT_DIR)], check=False)
        else:
            subprocess.run(["xdg-open", str(OUTPUT_DIR)], check=False)

    def reset_transaction_history(self) -> None:
        approved = messagebox.askyesno(
            "Reset transaction history?",
            "This deletes all imported statements and generated Excel output.\n\n"
            "Your saved merchant/category rules and custom categories will be kept. Source PDFs will not be deleted.",
            icon="warning",
            parent=self,
        )
        if not approved:
            return
        self.store.clear_transaction_history()
        if OUTPUT_DIR.exists():
            shutil.rmtree(OUTPUT_DIR)
        OUTPUT_DIR.mkdir(exist_ok=True)
        self.status.set("Transaction history and Excel output reset. Category rules were kept.")
        messagebox.showinfo("History reset", "Imported transaction history and generated Excel output were removed. Category rules were kept.", parent=self)

    def reset_all(self) -> None:
        approved = messagebox.askyesno(
            "Reset ALL data?",
            "This deletes the app's saved imports and merchant rules, and removes all files in the output folder.\n\n"
            "Your original PDF statements will not be deleted. This cannot be undone.",
            icon="warning",
            parent=self,
        )
        if not approved:
            return
        self.store.close()
        # SQLite can create these sidecar files when journal mode is enabled.
        for database_file in (DATABASE, Path(f"{DATABASE}-wal"), Path(f"{DATABASE}-shm")):
            if database_file.exists():
                database_file.unlink()
        if OUTPUT_DIR.exists():
            shutil.rmtree(OUTPUT_DIR)
        self.store = Store()
        OUTPUT_DIR.mkdir(exist_ok=True)
        self.status.set("Database and output folder reset. Select or re-import statements to test again.")
        messagebox.showinfo("Reset complete", "The local database and generated Excel output were removed.", parent=self)

    def on_close(self) -> None:
        self.store.close()
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
